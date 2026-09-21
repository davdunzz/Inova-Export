from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GREEN="00843D"
LIGHT="E7F4EC"

def export_xlsx(path, rows, fields):
    wb=Workbook()
    ws=wb.active
    ws.title="Dati"
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:{get_column_letter(len(fields))}{max(1,len(rows)+1)}"

    for col,field in enumerate(fields,1):
        c=ws.cell(1,col,field)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor=GREEN)
        c.alignment=Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[1].height=25

    thin=Side(style="thin",color="D5E0D9")
    for r_idx,row in enumerate(rows,2):
        for c_idx,field in enumerate(fields,1):
            c=ws.cell(r_idx,c_idx,row.get(field,""))
            c.alignment=Alignment(vertical="center")
            c.border=Border(bottom=thin)
            if r_idx % 2 == 0:
                c.fill=PatternFill("solid",fgColor="F5FAF7")

    for i,field in enumerate(fields,1):
        maxlen=len(field)
        for row in rows[:1000]:
            maxlen=max(maxlen,len(str(row.get(field,""))))
        ws.column_dimensions[get_column_letter(i)].width=min(max(12,maxlen+2),45)

    wb.save(path)
